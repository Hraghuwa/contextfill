from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = 'Security_Vulnerability_Audit_and_Remediation_Report.docx'
REVIEWED_COMMIT = '24cbb92262e844cd7f7e6a61a8f23807c9e6c185'
BLACK = '000000'
NAVY = '17365D'
BLUE_TINT = 'D9EAF7'
GRAY = 'D9D9D9'
PALE_GRAY = 'F4F6F8'


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)


def border(cell, color=GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = qn(f'w:{edge}')
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f'w:{edge}')
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')
        el.set(qn('w:color'), color)


def set_cell(cell, text, bold=False, color=BLACK, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.space_before = Pt(2)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    border(cell)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.autofit = False
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        if widths:
            cell.width = Inches(widths[idx])
        shade(cell, NAVY)
        set_cell(cell, header, bold=True, color='FFFFFF', align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            if widths:
                cells[idx].width = Inches(widths[idx])
            if row_index % 2:
                shade(cells[idx], PALE_GRAY)
            set_cell(cells[idx], value, align=WD_ALIGN_PARAGRAPH.CENTER if idx in (0, 2, 5) else WD_ALIGN_PARAGRAPH.LEFT)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_char1 = OxmlElement('w:fldChar')
    fld_char1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    fld_char2 = OxmlElement('w:fldChar')
    fld_char2.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_heading(doc, text, level=1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.color.rgb = RGBColor(0, 0, 0)
        run.font.name = 'Aptos Display' if level == 1 else 'Aptos'
    return paragraph


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.12
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)

    for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
        style = styles[name]
        style.font.name = 'Aptos Display' if name != 'Heading 2' else 'Aptos'
        style._element.rPr.rFonts.set(qn('w:ascii'), style.font.name)
        style._element.rPr.rFonts.set(qn('w:hAnsi'), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run('ContextFill Security Assessment')
    header_run.font.name = 'Aptos'
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run('Confidential assessment report | Page ')
    footer_run.font.name = 'Aptos'
    footer_run.font.size = Pt(8)
    add_field(footer, 'PAGE')


def add_cover(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(115)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title = p.add_run('Security Vulnerability Assessment and Remediation Report')
    title.bold = True
    title.font.name = 'Aptos Display'
    title._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display')
    title._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    title.font.size = Pt(27)
    title.font.color.rgb = RGBColor(0, 0, 0)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(16)
    subtitle.paragraph_format.space_after = Pt(36)
    run = subtitle.add_run('ContextFill Browser Extension')
    run.font.name = 'Aptos'
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(0, 0, 0)

    metadata = [
        ('Assessment date', 'October 2, 2026'),
        ('Assessment type', 'Application Security Code Review'),
        ('Prepared using', 'Codex-assisted Security Review'),
        ('Repository', 'ContextFill local repository'),
        ('Version reviewed', REVIEWED_COMMIT),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.autofit = False
    for label, value in metadata:
        cells = table.add_row().cells
        cells[0].width = Inches(1.85)
        cells[1].width = Inches(4.6)
        shade(cells[0], BLUE_TINT)
        set_cell(cells[0], label, bold=True)
        set_cell(cells[1], value)
    doc.add_paragraph()
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(45)
    run = note.add_run('Security vulnerabilities identified during the Codex-assisted security review were analyzed, remediated, and validated through code review and security testing.')
    run.italic = True
    run.font.size = Pt(10)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, 'Contents')
    contents = [
        '1 Executive Summary', '2 Application Overview', '3 Assessment Scope', '4 Security Methodology',
        '5 Vulnerability Summary', '6 Detailed Vulnerability Findings', '7 Remediation Details',
        '8 Security Testing Results', '9 Before and After Security Posture', '10 Remaining Risks',
        '11 Security Recommendations', '12 Conclusion',
    ]
    for item in contents:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.add_run(item)
    doc.add_page_break()


def add_report(doc):
    add_heading(doc, '1 Executive Summary')
    add_body(doc, 'This assessment reviewed the ContextFill Manifest V3 browser extension, which reads user-selected job-application forms and proposes field values from a locally stored profile and resume. The assessment focused on extension privileges, content-script and service-worker boundaries, AI-provider traffic, resume processing, configuration, and the existing automated workflow.')
    add_body(doc, 'Three confirmed security weaknesses were identified and remediated: persistent all-sites host access, excessive untrusted page data in AI requests combined with incomplete model-output controls, and an unrestricted Ollama endpoint with unbounded resume processing. No critical or high severity finding was confirmed. Post-remediation validation included syntax checks, Python compilation, diff validation, a Playwright end-to-end workflow, and a second targeted code review.')
    add_body(doc, 'The post-remediation posture is improved through least-privilege host permissions, data minimization, sensitive-field exclusion, strict output targeting, loopback-only local-model access, and bounded PDF processing. The review does not assert that the application is immune to future defects or hostile page behavior.')

    add_heading(doc, '2 Application Overview')
    add_body(doc, 'ContextFill is a Chrome and Safari extension that uses an explicit user action to inspect a current form, fill straightforward profile fields, and ask an AI provider to map the remaining fields. Users review populated values and submit the application themselves. The extension also stores a local application tracker and may attach a stored resume to recognized resume-upload fields.')
    add_table(doc, ['Component', 'Role', 'Trust boundary'], [
        ['Popup and options pages', 'Collect profile, model configuration, and user answers', 'Trusted extension UI; stores sensitive user data'],
        ['Service worker', 'Coordinates scans, model mapping, and filling', 'Privileged extension boundary'],
        ['Content script', 'Reads and updates active-tab form fields', 'Untrusted page DOM boundary'],
        ['Groq Claude Ollama', 'Generate structured field mappings', 'External provider or local loopback service'],
        ['PDF.js', 'Extracts local resume text', 'Untrusted uploaded document boundary'],
    ], [1.35, 2.45, 2.6])
    add_body(doc, 'The extension has no application-owned backend, database, login flow, server-side authorization layer, or background job system. Its sensitive state resides in browser extension storage. Its main data flow is profile and selected form metadata to the chosen model, structured suggestions back to the service worker, then visible form updates in the active tab.')

    add_heading(doc, '3 Assessment Scope')
    add_table(doc, ['Scope item', 'Reviewed evidence'], [
        ['Repository and branch', 'Local ContextFill repository on main'],
        ['Commit', REVIEWED_COMMIT],
        ['Source components', 'manifest.json, background.js, content.js, llm.js, options.js, popup.js, tracker.js'],
        ['Configuration and build', 'Manifest, Safari build script, repository history, bundled PDF.js metadata'],
        ['Tests', 'e2e_test.py, test-form.html, syntax and compilation checks'],
    ], [1.8, 4.6])

    add_heading(doc, '4 Security Methodology')
    for item in [
        'Repository reconnaissance and architecture review to establish data flows and trust boundaries.',
        'Source-code review for OWASP Top 10 and OWASP API Security themes, including authorization, injection, data handling, and configuration.',
        'Focused review of browser-extension permissions, service-worker messaging, content-script DOM access, local storage, and model-provider requests.',
        'Dependency and Git-history review for bundled component versions and exposed credentials.',
        'Remediation, security regression testing, and a second review of changed controls.',
    ]:
        add_bullet(doc, item)

    add_heading(doc, '5 Vulnerability Summary')
    add_table(doc, ['Severity', 'Identified', 'Fixed', 'Remaining'], [
        ['Critical', '0', '0', '0'], ['High', '0', '0', '0'], ['Medium', '2', '2', '0'],
        ['Low', '1', '1', '0'], ['Informational', '0', '0', '0'],
    ], [1.6, 1.35, 1.35, 1.55])
    add_table(doc, ['ID', 'Vulnerability', 'Severity', 'CWE', 'Location', 'Status'], [
        ['SEC-001', 'Persistent all-sites host permission', 'Medium', 'CWE-250', 'manifest.json', 'Fixed'],
        ['SEC-002', 'Excessive model input and output control gap', 'Medium', 'CWE-200 CWE-20', 'content.js llm.js', 'Fixed'],
        ['SEC-003', 'Unsafe local endpoint and unbounded resume input', 'Low', 'CWE-918 CWE-400', 'llm.js options.js', 'Fixed'],
    ], [0.65, 1.85, 0.75, 0.95, 1.4, 0.65])

    add_heading(doc, '6 Detailed Vulnerability Findings')
    detailed = [
        ('SEC-001 Persistent all-sites host permission', 'Medium', 'CWE-250 Execution with Unnecessary Privileges', 'OWASP A05 Security Misconfiguration',
         'Affected component: manifest.json host_permissions.',
         'Before remediation, the extension retained <all_urls> host access. A later extension defect or compromise would therefore have had standing access across sites beyond the current user-initiated form. The manifest now limits permanent access to the named AI APIs and loopback Ollama hosts. Active-tab access supports the current form and same-origin flow. Chromium loaded the manifest and the full active-tab workflow passed.'),
        ('SEC-002 Excessive model input and output control gap', 'Medium', 'CWE-200 Exposure of Sensitive Information and CWE-20 Improper Input Validation', 'OWASP A04 Insecure Design and A05 Security Misconfiguration',
         'Affected components: content.js scan and llm.js mapFields.',
         'Before remediation, up to 4,000 characters of arbitrary page transcript accompanied every model request, and duplicate or unknown returned field ids were not filtered. Page text may contain unrelated sensitive data or prompt-like instructions. The scanner now sends page URL and title only; model content is explicitly designated untrusted; sensitive identity, financial, and authentication fields become user-entered questions; and only the first result for each scanned id is accepted. Automated assertions confirmed all of these controls.'),
        ('SEC-003 Unsafe local endpoint and unbounded resume input', 'Low', 'CWE-918 Server Side Request Forgery and CWE-400 Uncontrolled Resource Consumption', 'OWASP A10 Server Side Request Forgery and A04 Insecure Design',
         'Affected components: llm.js and options.js.',
         'Before remediation, the configured Ollama URL was interpolated into requests and resumes had no explicit size, page, or extracted-text caps. The local model configuration now permits only loopback HTTP origins without credentials, paths, queries, or fragments. Resume parsing enforces 10 MB, 50-page, and 500,000-character limits and disables PDF scripting. The browser test confirmed remote endpoint rejection and normal local workflow behavior.'),
    ]
    for title, severity, cwe, owasp, component, text in detailed:
        add_heading(doc, title, level=2)
        add_body(doc, f'Severity: {severity}. {cwe}. {owasp}. ')
        add_body(doc, component)
        add_body(doc, text)

    add_heading(doc, '7 Remediation Details')
    add_table(doc, ['Finding', 'Root cause', 'Code change', 'Security control', 'Validation'], [
        ['SEC-001', 'Broad persistent host access', 'Replaced <all_urls> with five named hosts', 'Least privilege and active-tab scope', 'Chromium workflow passed'],
        ['SEC-002', 'Arbitrary transcript and unconstrained ids', 'Removed transcript; sensitive classifier; id allowlist', 'Data minimization and input validation', 'Targeted end-to-end assertions passed'],
        ['SEC-003', 'Arbitrary configured URL and no limits', 'Loopback validator; PDF limits; scripting off', 'Endpoint allowlist and resource bounds', 'Remote URL rejection passed'],
    ], [0.6, 1.35, 1.85, 1.55, 1.25])
    add_body(doc, 'The fixes are deliberately local to the existing extension architecture. They do not suppress warnings or disable checks, and the end-to-end workflow continues to require user review and user submission.')

    add_heading(doc, '8 Security Testing Results')
    add_table(doc, ['Test or scan', 'Result', 'Evidence'], [
        ['JavaScript syntax', 'Passed', 'node --check on all extension JavaScript files'],
        ['Python compilation', 'Passed', 'python3 -m py_compile on test utilities'],
        ['Diff validation', 'Passed', 'git diff --check returned no errors'],
        ['Playwright E2E', 'Passed', 'python3 e2e_test.py README.md returned PASS'],
        ['Ollama policy', 'Passed', 'Loopback accepted; remote endpoint rejected'],
        ['Sensitive fields', 'Passed', 'National ID converted to user-entered question'],
        ['Model mapping', 'Passed', 'Duplicate and unknown ids discarded'],
    ], [1.35, 0.8, 4.45])
    add_body(doc, 'No live Groq or Claude credentials were used. The end-to-end suite stubs those providers. A local Ollama model was available and its optional live check succeeded. Safari runtime validation was not executed because it requires macOS/Xcode packaging and manual extension enablement.')

    add_heading(doc, '9 Before and After Security Posture')
    add_table(doc, ['Area', 'Before', 'After'], [
        ['Host access', 'Persistent access to all URL origins', 'Named provider and loopback hosts; active-tab form access'],
        ['AI request context', 'Page title, URL, and arbitrary body transcript', 'Page title and URL plus selected field metadata'],
        ['Sensitive fields', 'Could be sent to model and filled', 'User-entered prompt; omitted from model mapping'],
        ['Model responses', 'Ids not deduplicated or allowlisted', 'First known scanned id only'],
        ['Resume parsing', 'No explicit limits; scripting default', 'Size page text limits and scripting disabled'],
    ], [1.25, 2.65, 2.7])

    add_heading(doc, '10 Remaining Risks')
    add_bullet(doc, 'Remote providers receive the profile, resume text, and form metadata needed for mapping when Groq or Claude is selected. This is a product data-sharing decision, not eliminated by the remediation.')
    add_bullet(doc, 'Label-based sensitive-field recognition reduces exposure but cannot replace user review of a hostile or unusual application page.')
    add_bullet(doc, 'Browser extension storage is isolated from normal web pages but is not application-level encrypted storage.')
    add_bullet(doc, 'Safari packaging and production browser compatibility were not manually penetration-tested during this code review.')

    add_heading(doc, '11 Security Recommendations')
    add_bullet(doc, 'Run the end-to-end and targeted security checks before releases that change permissions, model providers, or form-filling behavior.')
    add_bullet(doc, 'Track PDF.js releases and advisories; the bundled 6.3.289 version was checked against CVE-2026-16633 and is newer than its listed patched version 6.2.108.')
    add_bullet(doc, 'Keep API keys in trusted browser profiles and document the provider data-sharing boundary prominently for users.')
    add_bullet(doc, 'Perform manual hostile-page and Safari extension testing before a public production release.')
    add_bullet(doc, 'Consider a user-visible confirmation for unusually sensitive or long free-text fields as the product expands.')

    add_heading(doc, '12 Conclusion')
    add_body(doc, 'The review examined the extension manifest, extension pages, service worker, content script, model integration, local resume handling, dependency metadata, history, and automated tests. Three confirmed weaknesses were identified and remediated. Validation demonstrates that the revised controls work in the Chromium end-to-end flow and that no new confirmed issue was found in the second focused review. Remaining risks are documented above and should guide release testing and future security work.')


def main():
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_contents(doc)
    add_report(doc)
    doc.core_properties.title = 'Security Vulnerability Assessment and Remediation Report'
    doc.core_properties.subject = 'ContextFill application security code review'
    doc.core_properties.author = 'Codex-assisted Security Review'
    doc.core_properties.comments = 'Prepared from the actual repository audit and validation results.'
    doc.save(OUTPUT)


if __name__ == '__main__':
    main()
